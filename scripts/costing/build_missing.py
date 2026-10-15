"""Build the 'ingredients not on the price list' workbook from costs.json (curated groups over the raw missing lines)."""
import json,re,os,collections,datetime
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=os.path.abspath(os.path.join(os.path.dirname(__file__),'..','..'))
C=json.load(open(f'{ROOT}/content/costing/costs.json')); SH=json.load(open(f'{ROOT}/content/costing/sheet.json'))
RV=json.load(open(f'{ROOT}/content/costing/review.json'))
# (ingredient, category, unit, regex on the raw line)  — order matters (first match wins)
G=[('Ginger Garlic Paste','Sauces & Condiments','g',r'ginger.?garlic paste|garlic.?ginger paste|\bgg paste'),
('Greek Yogurt','Dairy & Eggs','g',r'greek yog|epigamia|high protein.*yog|plain.*yogurt'),
('Hing (Asafoetida)','Spices & Seasonings','g',r'\bhing\b|asafoetida'),
('Cinnamon Powder','Spices & Seasonings','g',r'cinnamon powder|cinnamon tsp'),
('Monk Fruit Sweetener','Pantry Staples','g',r'monk ?fruit|monkfruit'),
('Green Chutney','Sauces & Condiments','g',r'green chutney'),
('Chilli Vinegar','Oils & Vinegars','ml',r'chilli vinegar|chili vinegar'),
("Bird's Eye Chilli",'Vegetable','g',r'bird.?s? eye'),
('Dry Red Chilli','Spices & Seasonings','g',r'dr(y|ied) red chil|byadgi|kashmiri (dried )?red chil|kashmiri red chil'),
('Vegetable Stock (cube / powder)','Pantry Staples','g',r'veg(etable|e)?.{0,6}stock|stock powder'),
('Chicken Stock (cube)','Pantry Staples','g',r'chicken.{0,8}stock'),
('Banana','Vegetable','pc',r'\bbananas?\b'),
('Instant Coffee Powder','Pantry Staples','g',r'coffee powder|vietnamese coffee|^\s*coffee|brewed coffee|espresso'),
('Kasuri Methi','Spices & Seasonings','g',r'kas[ou]r?[iu] methi|kasturi'),
('Lemon','Vegetable','pc',r'\blemon\b(?! ?grass| ?pepper)|lemon.s juice|juice of'),
('Lime','Vegetable','pc',r'\blime\b(?! leaf)'),
('Tabasco / Hot Sauce','Sauces & Condiments','ml',r'tobasco|tabasco|hot sauce'),
('Worcestershire Sauce','Sauces & Condiments','ml',r'\bwc sauce|worcester'),
('Chicken Sausage','Poultry','g',r'sausage|frankfurter|salami'),
('Protein Powder','Pantry Staples','g',r'protein (powder|plant)|fast&up|chocolate protein|scoop.*protein|caramel protein|protein powder'),
('Dulce de Leche','Sauces & Condiments','g',r'dulce'),('Biscoff Spread','Sauces & Condiments','g',r'biscoff'),('Marshmallows','Flour & Baking','g',r'ma(r)?shmallow'),
('Nutella / Chocolate Spread','Sauces & Condiments','g',r'nutella'),('Oreo Biscuits','Flour & Baking','g',r'oreo'),('Biscuits (digestive / plain)','Flour & Baking','g',r'^\s*(\d+ )?biscuits?'),
('Cannelloni Tubes','Pantry Staples','g',r'cannelloni'),('Atta Noodles','Pantry Staples','g',r'atta noodles'),('Maggi Noodles','Pantry Staples','pc',r'maggi noodles|bundles maggi|maggi (masala )?noodles'),
('Frozen Momos','Frozen Foods','pc',r'\bmomos\b'),('Cauliflower','Vegetable','g',r'cauliflower'),('French Beans','Vegetable','g',r'french beans'),('Strawberries','Vegetable','g',r'strawberr'),
('Bok Choy','Vegetable','g',r'bok ?choy|bokchoy'),('Orange','Vegetable','pc',r'orange zest|\boranges?\b'),('Pineapple Juice','Sauces & Condiments','ml',r'pineapple juice'),
('Peri Peri Masala','Spices & Seasonings','g',r'peri.?peri masala'),('Shawarma Spice Mix','Spices & Seasonings','g',r'shawarma spice'),
('Kuzhambu / Rasam / Milagai Podi','Spices & Seasonings','g',r'kuzhambu podi|rasam podi|milagai podi|malaga podi|podi\b'),
('Dan Dan Sauce','Sauces & Condiments','ml',r'dan dan sauce'),('Chinsu Sauce','Sauces & Condiments','ml',r'chinsu'),('Maggi Hot & Sweet Sauce','Sauces & Condiments','ml',r'maggi hot'),
('Boombayway Sauces (stir fry / sweet mustard)','Sauces & Condiments','ml',r'boombayway.*(stir fry|sweet mustard|lime)'),('Chilli Crisp','Sauces & Condiments','g',r'chilli crisp'),
('Tamarind (paste / ball)','Pantry Staples','g',r'tamarind'),('Coconut Yogurt','Dairy & Eggs','g',r'coconut yogurt'),('Almond / Oat Milk','Dairy & Eggs','ml',r'almond.{0,8}milk|oats? milk'),
('Black Tea','Pantry Staples','g',r'black tea'),('Brandy / Cognac','Oils & Vinegars','ml',r'brandy|cognac'),('Gouda Cheese','Dairy & Eggs','g',r'gouda'),('Gruyère','Dairy & Eggs','g',r'gruy'),
('Idli / Dosa Batter','Pantry Staples','g',r'idli dosa batter'),('Boondi','Pantry Staples','g',r'boondi'),('Roti / Tortilla Wraps (cooked)','Bread & Bakery','pc',r'^\s*(\d+ )?rotis?\b|tortilla wraps'),
('Frozen Chicken Meatballs','Frozen Foods','g',r'chicken meatballs'),('Lettuce (generic leaves)','Vegetable','g',r'lettuce'),('Sweetener — Jaggery / Stevia mix','Pantry Staples','g',r'sugar stevia jaggery|sweetener'),
('Cooking Spray','Oils & Vinegars','ml',r'cooking spray'),('Brown Bread Croutons','Bread & Bakery','g',r'croutons'),('Mini Oreo / Cookie topping','Flour & Baking','g',r'mini oreo'),
('Green Chilli Paste / Red Chilli Paste','Sauces & Condiments','g',r'red chilli paste|chilli paste'),('Whole Grain / English Mustard','Sauces & Condiments','g',r'grain mustard|english mustard|mustard ketchup'),
]
G=[(a,b,c,re.compile(d,re.I)) for a,b,c,d in G]
cnt=collections.defaultdict(lambda:dict(recipes=set(),lines=0,ex=None)); other=collections.defaultdict(lambda:dict(recipes=set(),lines=0,ex=None))
for slug,r in C.items():
    for l in r['detail']:
        if l['status'] not in('missing','suggest'): continue
        raw=l['raw']
        for a,cat,u,rx in G:
            if rx.search(raw): e=cnt[a]; e['recipes'].add(slug); e['lines']+=1; e['ex']=e['ex'] or raw; break
        else:
            e=other[l['name']]; e['recipes'].add(slug); e['lines']+=1; e['ex']=e['ex'] or raw
meta={a:(cat,u) for a,cat,u,_ in G}
wb=Workbook(); H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='3B3F8F'); Y=PatternFill('solid',fgColor='FFF2CC')
def head(ws,cols,w):
    ws.append(cols)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,x in enumerate(w): ws.column_dimensions[chr(65+i)].width=x
    ws.freeze_panes='A2'
ws=wb.active; ws.title='Not on price list'
head(ws,['Ingredient','Category','Brand','Store','Pack size','Unit','Pack price (₹)','₹ per unit','Default?','Yield %','Updated on','Flag / note','Used in # recipes','Example from a recipe'],[34,20,16,14,10,7,12,11,9,9,12,30,12,60])
n=2
for a,e in sorted(cnt.items(),key=lambda x:-len(x[1]['recipes'])):
    cat,u=meta[a]
    ws.append([a,cat,None,None,None,u,None,f'=IF(AND(ISNUMBER(E{n}),ISNUMBER(G{n})),G{n}/E{n},"")','Y',1.0,None,'',len(e['recipes']),e['ex'][:110]])
    for col in 'CDEG': ws[f'{col}{n}'].fill=Y
    ws[f'J{n}'].number_format='0%'; ws[f'H{n}'].number_format='0.000'; n+=1
ws.auto_filter.ref=f'A1:N{n-1}'
# decide tab
ws=wb.create_sheet('Which one do you mean')
head(ws,['Recipe says','Used in # recipes','My provisional pick','Other options on your list','Example','Your answer (type the Ingredient name)'],[26,12,28,70,52,34])
import importlib.util
spec=importlib.util.spec_from_file_location('pc',f'{ROOT}/scripts/costing/parse_costs.py'); pc=importlib.util.module_from_spec(spec); spec.loader.exec_module(pc)
amb={x['key']:x for x in RV if x['status']=='ambiguous'}
usage=collections.defaultdict(lambda:[set(),None,0])
for slug,r in C.items():
    for l in r['detail']:
        if l.get('via')=='suggested': u=usage[l['name']]; u[0].add(slug); u[1]=u[1] or l['raw']; u[2]=l['item']
rows=[]
for nm,(rs,ex,item) in usage.items():
    k=pc.canon_name(nm); alts=[p['item'] for p in SH['prices'] if set(pc.toks(k))<=set(pc.toks(p['item'])) and p['item']!=item]
    rows.append((nm,len(rs),item,', '.join(sorted(set(alts))[:8]),ex[:70]))
for key,x in amb.items():
    rows.append((key,x['recipes'],'(none yet)',', '.join(x['candidates']),x['examples'][0][:70]))
for r in sorted(rows,key=lambda r:-r[1]): ws.append(list(r)); ws.cell(ws.max_row,6).fill=Y
# units tab
ws=wb.create_sheet('Quantities to confirm')
head(ws,['Ingredient (priced as)','Problem','Used in # recipes','Example line','Provisional value used','Your value'],[34,34,12,60,22,16])
PROV={k:v for k,v in pc.PIECE_G}
for x in [y for y in RV if y['status']=='unit']:
    it,reason=(x['key'].split('|')+[''])[:2]
    prob={'noqty':'recipe gives no quantity','piece-weight':'need grams per piece'}.get(reason,reason.replace('unit-','unit "')+('"' if reason.startswith('unit-') else ''))
    pv=next((f'{g} g each' for k,g in pc.PIECE_G if k in it.lower()),'')
    ws.append([it,prob,x['recipes'],x['examples'][0][:80],pv,None]); ws.cell(ws.max_row,6).fill=Y
# leftovers
ws=wb.create_sheet('Unclear text (I skipped)')
head(ws,['Parsed text','Used in # recipes','Example line'],[50,12,70])
for k,e in sorted(other.items(),key=lambda x:-len(x[1]['recipes'])):
    if len(k)>1: ws.append([k,len(e['recipes']),e['ex'][:110]])
wb.save(f'{ROOT}/content/costing/Not-on-price-list.xlsx')
print('groups',len(cnt),'recipes touched',len({s for e in cnt.values() for s in e['recipes']}),'leftover texts',len(other))
print([(a,len(e['recipes'])) for a,e in sorted(cnt.items(),key=lambda x:-len(x[1]['recipes']))][:25])
