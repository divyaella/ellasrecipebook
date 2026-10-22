"""Workbook of everything still blocking the incomplete recipes (hidden recipes excluded)."""
import json,re,os,collections
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
ROOT=os.path.abspath(os.path.join(os.path.dirname(__file__),'..','..'))
C=json.load(open(f'{ROOT}/content/costing/costs.json')); SH=json.load(open(f'{ROOT}/content/costing/sheet.json'))
import sys; sys.path.insert(0,os.path.dirname(__file__)); from groups import G
have={p['item'].lower() for p in SH['prices']}
def group(raw):
    for a,cat,u,rx in G:
        if rx.search(raw): return a,cat,u
    return None,None,None
inc={k:v for k,v in C.items() if not v['complete']}
issues=collections.defaultdict(lambda:dict(type=None,recipes=set(),ex=None,cands=set(),grp=None,cat=None,unit=None,item=None,reason=None))
byrec=[]
for slug,v in inc.items():
    open_=[]
    for l in v['detail']:
        st=l['status']
        if st in('ok','vague','free'): continue
        if st in('missing','suggest'):
            a,cat,u=group(l['raw'])
            if a and a.lower() in have: continue           # already priced under the group name
            if not a and len(l['name'].split())>=4: continue
            key=('price',a or l['name'])
            e=issues[key]; e['type']='Needs a price / ingredient'; e['grp']=a; e['cat']=cat; e['unit']=u
        elif st=='ambiguous':
            key=('which',l['name']); e=issues[key]; e['type']='Which ingredient do you mean?'
        elif st=='unit':
            key=('weight',(l.get('item') or l['name'])+'|'+l.get('reason','')); e=issues[key]; e['type']='Weight / unit needed'; e['item']=l.get('item'); e['reason']=l.get('reason')
        else: continue
        e['recipes'].add(slug); e['ex']=e['ex'] or l['raw']; e['cands']|=set(l.get('candidates',[]))
        open_.append(l['raw'][:60])
    byrec.append((v['title'],v['serves'],len(open_),' || '.join(open_[:4])))
wb=Workbook(); H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='3B3F8F'); Y=PatternFill('solid',fgColor='FFF2CC')
def sheet(name,cols,widths,first=False):
    ws=wb.active if first else wb.create_sheet(name); ws.title=name; ws.append(cols)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
    ws.freeze_panes='A2'; return ws
ws=sheet('1 Prices to add',['Ingredient to add','Category','Brand','Store','Pack size','Unit','Pack price (₹)','Blocks # recipes','Example line','Note'],[34,20,22,12,10,7,12,10,60,34],first=True)
for (t,k),e in sorted([i for i in issues.items() if i[0][0]=='price'],key=lambda x:-len(x[1]['recipes'])):
    ws.append([e['grp'] or k.title(),e['cat'],None,'Blinkit',None,e['unit'],None,len(e['recipes']),e['ex'][:90],'' if e['grp'] else 'check the wording — not one of my standard groups'])
    for col in 'CEFG': ws[f'{col}{ws.max_row}'].fill=Y
ws=sheet('2 Which one',['Recipe says','Blocks # recipes','Options on your list','Example line','Your answer (ingredient name)'],[26,10,80,56,34])
for (t,k),e in sorted([i for i in issues.items() if i[0][0]=='which'],key=lambda x:-len(x[1]['recipes'])):
    ws.append([k,len(e['recipes']),', '.join(sorted(e['cands']))[:200],e['ex'][:80],None]); ws.cell(ws.max_row,5).fill=Y
ws=sheet('3 Weight needed',['Ingredient (as priced)','Problem','Blocks # recipes','Example line','Grams (or ml) each'],[34,26,10,60,18])
for (t,k),e in sorted([i for i in issues.items() if i[0][0]=='weight'],key=lambda x:-len(x[1]['recipes'])):
    ws.append([e['item'] or k.split('|')[0],(e['reason'] or '').replace('piece-weight','weight per piece needed').replace('unit-','unit not understood: '),len(e['recipes']),e['ex'][:80],None]); ws.cell(ws.max_row,5).fill=Y
ws=sheet('By recipe',['Recipe','Serves','Open lines','What is open'],[48,8,10,120])
for r in sorted(byrec,key=lambda r:r[2]): ws.append(list(r))
wb.save(f'{ROOT}/content/costing/Pending-issues.xlsx')
print({t:sum(1 for k in issues if k[0]==t) for t in('price','which','weight')},'incomplete recipes',len(inc))
